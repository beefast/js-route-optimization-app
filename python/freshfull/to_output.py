#!/usr/bin/env python3
"""
Convert a Route Optimization API response into the Freshfull output format.

Produces an Excel file matching the "Rezultate simulare.xlsx" schema:

    DataStartLivrare        — actual delivery visit start time
    DataEndLivrare          — actual delivery visit end time  (start + 7 min)
    Comanda                 — order ID
    DataPlasareComanda      — order placement datetime (from original request)
    Sofer                   — driver label (vehicle label used as driver proxy)
    IntervalClient          — customer-requested delivery time window
    DataLivrareSolicitataClient — end of customer's requested window
    TipVehicul              — Iveco / Peugeot / Renault
    TipProgramSofer         — 8h / 10h / 12h (inferred from routeDurationLimit)
    Cursa                   — unique route identifier
    DataStartCursa          — route start time (vehicle departs depot)
    DataEndCursa            — route end time (vehicle returns to depot)

Usage:
    python to_output.py \\
        --request  data/freshfull/request_sample.json \\
        --response data/freshfull/response_sample.json \\
        --output   data/freshfull/rezultate_sample.xlsx
"""

import argparse
import json
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

UTC_OFFSET_HOURS = 2   # Romania EET; match the value used in transform.py
DELIVERY_DURATION_MIN = 7

HEADERS = [
    "DataStartLivrare",
    "DataEndLivrare",
    "Comanda",
    "DataPlasareComanda",
    "Sofer",
    "IntervalClient",
    "DataLivrareSolicitataClient",
    "TipVehicul",
    "TipProgramSofer",
    "Cursa",
    "DataStartCursa",
    "DataEndCursa",
]

HEADER_FILL  = PatternFill("solid", fgColor="D9E1F2")
HEADER_FONT  = Font(bold=True, name="Calibri")
BODY_FONT    = Font(name="Calibri")
CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
LEFT_ALIGN   = Alignment(horizontal="left",   vertical="center")

_TZ = timezone(timedelta(hours=UTC_OFFSET_HOURS))


def parse_ts(s: str | None) -> datetime | None:
    """Parse RFC 3339 / ISO 8601 string to UTC-aware datetime."""
    if not s:
        return None
    s = s.replace("Z", "+00:00")
    return datetime.fromisoformat(s).astimezone(timezone.utc)


def to_local(dt: datetime | None) -> datetime | None:
    """Convert UTC datetime to Romania local time."""
    if dt is None:
        return None
    return dt.astimezone(_TZ)


def fmt(dt: datetime | None) -> str:
    """Format datetime for Excel cell (local time, human-readable)."""
    if dt is None:
        return ""
    return dt.strftime("%Y-%m-%d %H:%M")


def vehicle_type(label: str) -> str:
    prefix = label.split("-")[0].upper()
    return {"IV": "Iveco", "PG": "Peugeot", "RN": "Renault"}.get(prefix, label)


def shift_type_from_request(vehicle_dict: dict) -> str:
    """Read routeDurationLimit from the request vehicle to determine shift type."""
    limit = vehicle_dict.get("routeDurationLimit", {})
    dur_str = limit.get("maxDuration", "36000s")
    hours = int(dur_str.rstrip("s")) // 3600
    if hours <= 8:
        return "8h"
    if hours <= 10:
        return "10h"
    return "12h"


def build_output_rows(request: dict, response: dict) -> list[dict]:
    """Join API response routes/visits with the original request to produce output rows."""
    shipments = request["model"]["shipments"]
    vehicles  = request["model"]["vehicles"]
    rows = []

    for route_idx, route in enumerate(response.get("routes", [])):
        v_idx     = route.get("vehicleIndex", 0)
        vehicle   = vehicles[v_idx]
        v_label   = vehicle.get("label", f"V-{v_idx:03d}")
        v_type    = vehicle_type(v_label)
        shift     = shift_type_from_request(vehicle)

        route_start = to_local(parse_ts(route.get("vehicleStartTime")))
        route_end   = to_local(parse_ts(route.get("vehicleEndTime")))
        cursa_id    = (
            f"{v_label}-"
            + (route_start.strftime("%m%d-%H%M") if route_start else f"R{route_idx:04d}")
        )

        for visit in route.get("visits", []):
            s_idx    = visit.get("shipmentIndex", 0)
            shipment = shipments[s_idx]
            order_id = shipment.get("label", str(s_idx))

            # Visit timing
            visit_start = to_local(parse_ts(visit.get("startTime")))
            visit_end   = (visit_start + timedelta(minutes=DELIVERY_DURATION_MIN)) if visit_start else None

            # Customer time window (from the request's delivery timeWindow)
            delivery_visits = shipment.get("deliveries", [])
            tw_start_local = tw_end_local = None
            if delivery_visits:
                tw = (delivery_visits[0].get("timeWindows") or [{}])[0]
                tw_start_local = to_local(parse_ts(tw.get("startTime")))
                tw_end_local   = to_local(parse_ts(tw.get("endTime")))

            interval_str = (
                f"{tw_start_local.strftime('%H:%M')} - {tw_end_local.strftime('%H:%M')}"
                if tw_start_local and tw_end_local else ""
            )

            # Order placement (stored in request as first timeWindow startTime adjusted
            # by prep time — we don't have exact placement in the response, so we
            # reconstruct: placement ≈ effective_window_start - 125 min)
            placement = None
            if tw_start_local:
                placement = tw_start_local - timedelta(minutes=125)

            rows.append({
                "DataStartLivrare":           fmt(visit_start),
                "DataEndLivrare":             fmt(visit_end),
                "Comanda":                    order_id,
                "DataPlasareComanda":         fmt(placement),
                "Sofer":                      v_label,
                "IntervalClient":             interval_str,
                "DataLivrareSolicitataClient":fmt(tw_end_local),
                "TipVehicul":                 v_type,
                "TipProgramSofer":            shift,
                "Cursa":                      cursa_id,
                "DataStartCursa":             fmt(route_start),
                "DataEndCursa":               fmt(route_end),
            })

    return rows


def write_excel(rows: list[dict], output_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rezultate"
    ws.freeze_panes = "A2"

    # Header row
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = CENTER_ALIGN

    # Data rows
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, key in enumerate(HEADERS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row.get(key, ""))
            cell.font      = BODY_FONT
            cell.alignment = LEFT_ALIGN

    # Auto-size columns
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 28)

    wb.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Convert Route Optimization API response to Freshfull output Excel."
    )
    parser.add_argument("--request",  required=True, help="Request JSON (from transform.py)")
    parser.add_argument("--response", required=True, help="Response JSON (from run_api.py)")
    parser.add_argument("--output",   required=True, help="Output Excel file path (.xlsx)")
    args = parser.parse_args()

    request  = json.loads(Path(args.request).read_text())
    response = json.loads(Path(args.response).read_text())

    rows = build_output_rows(request, response)

    n_served  = len(rows)
    n_skipped = len(response.get("skippedShipments", []))
    print(f"[to_output] {n_served} deliveries assigned, {n_skipped} skipped")

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    write_excel(rows, output_path)
    print(f"[to_output] Saved to {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
