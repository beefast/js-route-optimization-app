#!/usr/bin/env python3
"""
Transform Freshfull delivery data (Excel) into a Google Route Optimization API request.

Reads "Centralizator 23-29 martie.xlsx", filters by date, and produces a JSON request
compatible with the Route Optimization API (routeoptimization.googleapis.com).

Usage:
    # Small sample for cheap testing
    python transform.py \\
        --input "Centralizator 23-29 martie.xlsx" \\
        --date 2026-03-23 \\
        --sample 50 \\
        --vehicles 20 \\
        --output data/freshfull/request_sample.json

    # Full day run
    python transform.py \\
        --input "Centralizator 23-29 martie.xlsx" \\
        --date 2026-03-23 \\
        --high-quality \\
        --timeout 600 \\
        --output data/freshfull/request_2026-03-23.json
"""

import argparse
import json
import math
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl

# ──────────────────────────────────────────────
# Geography & timing constants
# ──────────────────────────────────────────────

DEPOT: dict = {"latitude": 44.52655539645606, "longitude": 26.18054083552211}
UTC_OFFSET_HOURS = 2              # Romania EET (UTC+2); switch to 3 after DST on Mar 29
DELIVERY_DURATION_S = 7 * 60     # 7 min per customer stop
PREP_MINUTES = 125                # min advance booking window (order → earliest delivery)
GLOBAL_START_HOUR = 6            # 06:00 local — earliest vehicle departure
GLOBAL_END_HOUR = 23             # 23:00 local — latest allowed time in model

# ──────────────────────────────────────────────
# Vehicle fleet specification
# ──────────────────────────────────────────────
#
# Load dimensions:
#   ambient          — ambient-temperature boxes (Peugeot has a separate compartment cap)
#   chilled          — chilled boxes (Peugeot has a separate compartment cap)
#   frozen           — frozen bags (all vehicles share 14-bag limit)
#   ambient_chilled  — combined ambient+chilled (Iveco/Renault have a single shared hold)
#
# Setting a limit to 9999 effectively makes it non-binding.

FLEET = [
    {
        "type": "Iveco",
        "prefix": "IV",
        "count": 70,
        "cost_fixed": 120.0,
        "cost_per_hour": 60.0,
        "cost_per_km": 2.0,
        "shift_hours": 10,
        "limits": {
            "ambient":         {"maxLoad": "9999"},
            "chilled":         {"maxLoad": "9999"},
            "frozen":          {"maxLoad": "14"},
            "ambient_chilled": {"maxLoad": "120"},
        },
    },
    {
        "type": "Peugeot",
        "prefix": "PG",
        "count": 85,
        "cost_fixed": 100.0,
        "cost_per_hour": 50.0,
        "cost_per_km": 1.5,
        "shift_hours": 10,
        "limits": {
            "ambient":         {"maxLoad": "48"},
            "chilled":         {"maxLoad": "39"},
            "frozen":          {"maxLoad": "14"},
            "ambient_chilled": {"maxLoad": "9999"},  # Peugeot has separate compartments
        },
    },
    {
        "type": "Renault",
        "prefix": "RN",
        "count": 12,
        "cost_fixed": 150.0,
        "cost_per_hour": 70.0,
        "cost_per_km": 2.5,
        "shift_hours": 10,
        "limits": {
            "ambient":         {"maxLoad": "9999"},
            "chilled":         {"maxLoad": "9999"},
            "frozen":          {"maxLoad": "14"},
            "ambient_chilled": {"maxLoad": "144"},
        },
    },
]

PENALTY_COST = 10_000        # cost of skipping a shipment — kept high to force coverage
LOADING_DOCKS = 28           # max vehicles loading simultaneously
BATCH_WINDOW_MIN = 30        # each dock batch gets a 30-min start window slot
BREAK_EARLIEST_HOUR = 11     # break window: 11:00 local
BREAK_LATEST_HOUR = 16       # break window: 16:00 local
BREAK_DURATION_S = 3600      # 1-hour mandatory break


# ──────────────────────────────────────────────
# Time helpers
# ──────────────────────────────────────────────

def _tz() -> timezone:
    return timezone(timedelta(hours=UTC_OFFSET_HOURS))


def excel_to_utc(value) -> datetime:
    """Convert an Excel cell value (float serial or Python datetime) to UTC-aware datetime."""
    if isinstance(value, datetime):
        dt = value.replace(tzinfo=_tz()) if value.tzinfo is None else value
    elif isinstance(value, (int, float)):
        dt = (datetime(1899, 12, 30) + timedelta(days=float(value))).replace(tzinfo=_tz())
    else:
        raise ValueError(f"Cannot convert {type(value).__name__!r} to datetime: {value!r}")
    return dt.astimezone(timezone.utc)


def ts(dt: datetime) -> str:
    """Format a UTC-aware datetime as RFC 3339 Zulu string (used in all API time fields)."""
    utc = dt.astimezone(timezone.utc)
    return utc.strftime("%Y-%m-%dT%H:%M:%SZ")


def local(date, hour: int, minute: int = 0) -> datetime:
    """Build a UTC-aware datetime from a local date + hour:minute."""
    return datetime(date.year, date.month, date.day, hour, minute, tzinfo=_tz())


# ──────────────────────────────────────────────
# Vehicle builder
# ──────────────────────────────────────────────

def build_vehicles(date, max_count: int | None = None) -> tuple[list[dict], dict[str, list[int]]]:
    """
    Build the vehicles list for the API request and return a type→indices mapping.

    Loading-dock constraint: stagger vehicle start windows in batches of LOADING_DOCKS.
    Each batch gets an exclusive BATCH_WINDOW_MIN-minute slot starting from GLOBAL_START_HOUR.
    This ensures at most LOADING_DOCKS vehicles start loading simultaneously.

    Returns:
        vehicles:   list of Vehicle dicts ready for the API request
        type_index: {"Iveco": [0,1,...], "Peugeot": [70,...], "Renault": [155,...]}
    """
    total_fleet = sum(spec["count"] for spec in FLEET)

    vehicles: list[dict] = []
    type_index: dict[str, list[int]] = {spec["type"]: [] for spec in FLEET}

    for spec in FLEET:
        count = spec["count"]
        if max_count is not None:
            # Scale each type proportionally, minimum 1
            count = max(1, round(spec["count"] / total_fleet * max_count))

        for i in range(count):
            if max_count is not None and len(vehicles) >= max_count:
                break

            vid = len(vehicles)
            type_index[spec["type"]].append(vid)

            # Stagger start window based on dock batch
            batch = vid // LOADING_DOCKS
            batch_start_min = batch * BATCH_WINDOW_MIN
            slot_start = local(date, GLOBAL_START_HOUR) + timedelta(minutes=batch_start_min)
            slot_end = slot_start + timedelta(minutes=BATCH_WINDOW_MIN - 1)

            # Cap slot within global bounds
            global_end = local(date, GLOBAL_END_HOUR)
            if slot_start >= global_end:
                slot_start = local(date, GLOBAL_START_HOUR)
                slot_end = slot_start + timedelta(minutes=BATCH_WINDOW_MIN - 1)

            vehicle: dict = {
                "label": f"{spec['prefix']}-{i + 1:03d}",
                "travelMode": "DRIVING",
                "startLocation": DEPOT,
                "endLocation": DEPOT,
                "startTimeWindows": [{"startTime": ts(slot_start), "endTime": ts(slot_end)}],
                "fixedCost": spec["cost_fixed"],
                "costPerHour": spec["cost_per_hour"],
                "costPerKilometer": spec["cost_per_km"],
                "usedIfRouteIsEmpty": False,
                "loadLimits": spec["limits"],
                # Enforce total shift length (routing time + break must fit within shift_hours)
                "routeDurationLimit": {"maxDuration": f"{spec['shift_hours'] * 3600}s"},
                "breakRule": {
                    "breakRequests": [{
                        "earliestStartTime": ts(local(date, BREAK_EARLIEST_HOUR)),
                        "latestStartTime":   ts(local(date, BREAK_LATEST_HOUR)),
                        "minDuration":       f"{BREAK_DURATION_S}s",
                    }]
                },
            }
            vehicles.append(vehicle)

    return vehicles, type_index


# ──────────────────────────────────────────────
# Shipment builder
# ──────────────────────────────────────────────

def _is_restricted(vehicle_info: str) -> bool:
    """Return True if this order is restricted to Peugeot-only vehicles."""
    if not vehicle_info:
        return False
    lower = str(vehicle_info).lower()
    # "toate tipurile" → all vehicle types allowed
    if "toate" in lower:
        return False
    # Any other value is treated as a vehicle restriction
    return True


def build_shipments(rows: list[dict], vehicles: list[dict], type_index: dict[str, list[int]], date) -> list[dict]:
    """Build the shipments list from Excel rows."""
    peugeot_indices = type_index.get("Peugeot", [])

    shipments = []
    for row in rows:
        order_id = str(row["IdComanda"]).strip()
        lat = float(row["Coordonate.geo_lat"])
        lon = float(row["Coordonate.geo_long"])

        # Load demands — ceil to nearest whole box/bag
        ambient = math.ceil(float(row["CapacitatePlanificataZonaAmbient"] or 0))
        chilled = math.ceil(float(row["CapacitatePlanificataZonaChilled"] or 0))
        frozen  = math.ceil(float(row["CapacitatePlanificataZonaFrozen"]  or 0))

        # Delivery time window from pre-computed interval columns
        tw_start = excel_to_utc(row["DataStartLivrareIntervalClient"])
        tw_end   = excel_to_utc(row["DataEndLivrareIntervalClient"])

        # Earliest possible delivery = order placement + 125 min prep
        order_placed = excel_to_utc(row["Data_plasare_comanda"])
        earliest = order_placed + timedelta(minutes=PREP_MINUTES)
        effective_start = max(tw_start, earliest)

        # Guard: if prep pushes past window end, clamp to window end - 1 min
        if effective_start >= tw_end:
            effective_start = tw_end - timedelta(minutes=1)

        shipment: dict = {
            "label": order_id,
            "penaltyCost": PENALTY_COST,
            "deliveries": [{
                "arrivalLocation": {"latitude": lat, "longitude": lon},
                "duration": f"{DELIVERY_DURATION_S}s",
                "timeWindows": [{
                    "startTime": ts(effective_start),
                    "endTime":   ts(tw_end),
                }],
            }],
            "loadDemands": {
                "ambient":         {"amount": str(ambient)},
                "chilled":         {"amount": str(chilled)},
                "frozen":          {"amount": str(frozen)},
                "ambient_chilled": {"amount": str(ambient + chilled)},
            },
        }

        # Zone restriction: only Peugeot for city-centre / residence zones
        vehicle_info = str(row.get("InfoTipMasinaLivrare") or "")
        if _is_restricted(vehicle_info) and peugeot_indices:
            shipment["allowedVehicleIndices"] = peugeot_indices

        shipments.append(shipment)

    return shipments


# ──────────────────────────────────────────────
# Full request builder
# ──────────────────────────────────────────────

def build_request(
    rows: list[dict],
    date,
    max_vehicles: int | None = None,
    high_quality: bool = False,
    timeout_s: int = 60,
) -> dict:
    vehicles, type_index = build_vehicles(date, max_count=max_vehicles)
    shipments = build_shipments(rows, vehicles, type_index, date)

    return {
        "model": {
            "shipments": shipments,
            "vehicles": vehicles,
            "globalStartTime": ts(local(date, GLOBAL_START_HOUR)),
            "globalEndTime":   ts(local(date, GLOBAL_END_HOUR)),
        },
        "searchMode": "RETURN_HIGH_QUALITY" if high_quality else "RETURN_FAST",
        "timeout": f"{timeout_s}s",
        "considerRoadTraffic": True,
        "label": f"Freshfull-{date.isoformat()}",
    }


# ──────────────────────────────────────────────
# Excel reader
# ──────────────────────────────────────────────

def read_excel(path: Path, target_date, sample: int | None = None) -> list[dict]:
    """
    Read the Centralizator Excel and return rows for target_date.
    Rows are matched by the date part of DataStartLivrareIntervalClient.
    If sample is set, takes an evenly-spaced subset across the day.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

    rows = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, raw))
        val = row.get("DataStartLivrareIntervalClient")
        if val is None:
            continue
        try:
            delivery_utc = excel_to_utc(val)
            local_date = delivery_utc.astimezone(_tz()).date()
            if local_date != target_date:
                continue
            rows.append(row)
        except (ValueError, TypeError):
            continue

    wb.close()

    if sample and len(rows) > sample:
        step = len(rows) / sample
        rows = [rows[int(i * step)] for i in range(sample)]

    return rows


# ──────────────────────────────────────────────
# CLI
# ──────────────────────────────────────────────

def main() -> int:
    parser = argparse.ArgumentParser(
        description="Convert Freshfull Excel data to a Route Optimization API request JSON."
    )
    parser.add_argument("--input",   required=True, help="Path to Centralizator Excel file")
    parser.add_argument("--date",    required=True, help="Date to process (YYYY-MM-DD)")
    parser.add_argument("--output",  required=True, help="Output JSON file path")
    parser.add_argument("--sample",  type=int, default=None, help="Limit to N orders (for cheap testing)")
    parser.add_argument("--vehicles",type=int, default=None, help="Limit total vehicle count")
    parser.add_argument("--high-quality", action="store_true", help="Use RETURN_HIGH_QUALITY mode (slower, better)")
    parser.add_argument("--timeout", type=int, default=60, help="Solver timeout in seconds (default: 60)")
    args = parser.parse_args()

    target_date = datetime.strptime(args.date, "%Y-%m-%d").date()
    input_path  = Path(args.input)
    output_path = Path(args.output)

    print(f"[transform] Reading {input_path} for {target_date}...")
    rows = read_excel(input_path, target_date, sample=args.sample)
    print(f"[transform] {len(rows)} orders loaded")

    if not rows:
        print("[transform] ERROR: no orders found for that date.")
        return 1

    request = build_request(
        rows,
        date=target_date,
        max_vehicles=args.vehicles,
        high_quality=args.high_quality,
        timeout_s=args.timeout,
    )

    n_v = len(request["model"]["vehicles"])
    n_s = len(request["model"]["shipments"])
    print(f"[transform] Built request: {n_s} shipments, {n_v} vehicles")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(request, indent=2, ensure_ascii=False))
    print(f"[transform] Saved to {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
